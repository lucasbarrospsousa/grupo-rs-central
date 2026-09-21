import argparse
import json
import os
import sys
from collections import Counter
from datetime import datetime


NAVY = "0B3154"
BLUE = "0B7FC1"
GREEN = "14A66B"
ORANGE = "F59E0B"
LIGHT = "F4F8FC"
GRID = "D8E4EF"
TEXT = "17263A"
MUTED = "65758B"


def clean(value, fallback="Não informado"):
    text = str(value or "").strip()
    return text if text else fallback


def product_rows(payload):
    result = []
    for item in payload.get("products", []):
        serial = clean(item.get("serial"), "-")
        result.append({
            "serial": serial,
            "plate": clean(item.get("plate"), "Sem placa"),
            "model": clean(item.get("model"), "Não informado"),
            "operator": clean(item.get("operator"), "Não informada"),
            "status": clean(item.get("status"), "Não informado"),
            "connectivity": clean(item.get("connectivity"), "Não consultado"),
            "updated": clean(item.get("updated"), "-"),
        })
    return result


def distributions(rows, key):
    return Counter(row[key] for row in rows)


def generate_xlsx(payload, output_path, logo_path):
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.table import Table, TableStyleInfo

    rows = product_rows(payload)
    wb = Workbook()
    ws = wb.active
    ws.title = "Relatório"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.outlinePr.summaryBelow = True
    ws.freeze_panes = "A15"

    widths = [18, 18, 22, 18, 18, 22]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + idx)].width = width

    if logo_path and os.path.exists(logo_path):
        image = Image(logo_path)
        image.width, image.height = 62, 62
        ws.add_image(image, "A1")

    ws.merge_cells("B1:F1")
    ws["B1"] = "GRUPO RS CENTRAL"
    ws["B1"].font = Font(name="Aptos Display", size=23, bold=True, color=NAVY)
    ws["B1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 34
    ws.merge_cells("B2:F2")
    ws["B2"] = "GESTÃO INTELIGENTE DE EQUIPAMENTOS"
    ws["B2"].font = Font(name="Aptos", size=9, bold=True, color=ORANGE)

    ws.merge_cells("A4:F4")
    ws["A4"] = payload.get("title", "Relatório de estoque")
    ws["A4"].font = Font(name="Aptos Display", size=20, bold=True, color=TEXT)
    ws.merge_cells("A5:F5")
    ws["A5"] = f"Filial: {payload.get('branch', '-')}  •  Situação: {payload.get('status', '-')}  •  Gerado em: {payload.get('generated_at', '-')}"
    ws["A5"].font = Font(name="Aptos", size=10, color=MUTED)

    total = len(rows)
    operators = distributions(rows, "operator")
    models = distributions(rows, "model")
    missing_chip = sum(1 for item in payload.get("products", []) if any(key in item for key in ("iccid", "chip_number", "numero_chip")) and not str(item.get("iccid") or item.get("chip_number") or item.get("numero_chip") or "").strip())
    cards = [("TOTAL", total), ("VIVO", operators.get("Vivo", 0)), ("CLARO", operators.get("Claro", 0)), ("SEM CHIP", missing_chip)]
    for index, (label, value) in enumerate(cards):
        col = 1 + index
        cell = ws.cell(7, col)
        cell.value = label
        cell.font = Font(name="Aptos", size=9, bold=True, color=MUTED)
        cell.fill = PatternFill("solid", fgColor=LIGHT)
        cell.alignment = Alignment(horizontal="center")
        val = ws.cell(8, col)
        val.value = value
        val.font = Font(name="Aptos Display", size=18, bold=True, color=NAVY)
        val.fill = PatternFill("solid", fgColor=LIGHT)
        val.alignment = Alignment(horizontal="center")

    ws.merge_cells("A10:F10")
    ws["A10"] = "Distribuição por modelo"
    ws["A10"].font = Font(name="Aptos Display", size=13, bold=True, color=TEXT)
    model_text = "   •   ".join(f"{name}: {count}" for name, count in models.most_common()) or "Sem dados"
    ws.merge_cells("A11:F11")
    ws["A11"] = model_text
    ws["A11"].font = Font(name="Aptos", size=10, color=MUTED)

    headers = ["Número de série", "Identificação", "Modelo", "Operadora", "Situação", "Atualização"]
    for col, value in enumerate(headers, 1):
        cell = ws.cell(14, col, value)
        cell.font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[14].height = 25

    thin = Side(style="thin", color=GRID)
    for r_index, row in enumerate(rows, 15):
        values = [row["serial"], row["plate"], row["model"], row["operator"], row["status"], row["updated"]]
        for c_index, value in enumerate(values, 1):
            cell = ws.cell(r_index, c_index, value)
            cell.font = Font(name="Aptos", size=10, color=TEXT)
            cell.fill = PatternFill("solid", fgColor="FFFFFF" if r_index % 2 else LIGHT)
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(vertical="center")
        ws.row_dimensions[r_index].height = 22
    if rows:
        table = Table(displayName="TabelaEquipamentos", ref=f"A14:F{14 + len(rows)}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False, showLastColumn=False)
        ws.add_table(table)

    last = max(16, 15 + len(rows))
    ws.merge_cells(start_row=last, start_column=1, end_row=last, end_column=6)
    ws.cell(last, 1).value = "Grupo RS Central • Relatório gerado pelo sistema"
    ws.cell(last, 1).font = Font(name="Aptos", size=9, color=MUTED)
    ws.cell(last, 1).alignment = Alignment(horizontal="center")
    ws.print_title_rows = "14:14"
    ws.print_area = f"A1:F{last}"
    ws.oddFooter.center.text = "Grupo RS Central • Página &P de &N"
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    return {"ok": True, "path": output_path, "format": "xlsx", "rows": len(rows), "sheets": len(wb.sheetnames)}


def generate_pdf(payload, output_path, logo_path):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT, TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.graphics.shapes import Circle, Drawing, Line, Path, Rect
    from reportlab.platypus import Image, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    rows = product_rows(payload)
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="ReportTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=18, leading=22, textColor=colors.HexColor("#17263A"), alignment=TA_LEFT, spaceAfter=3))
    styles.add(ParagraphStyle(name="Meta", parent=styles["Normal"], fontName="Helvetica", fontSize=8.5, textColor=colors.HexColor("#65758B"), leading=11))
    styles.add(ParagraphStyle(name="Section", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=11.5, textColor=colors.HexColor("#17263A"), spaceBefore=7, spaceAfter=5))
    styles.add(ParagraphStyle(name="Footer", parent=styles["Normal"], fontName="Helvetica", fontSize=7.5, textColor=colors.HexColor("#65758B")))

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor("#D8E4EF"))
        canvas.line(16 * mm, 13 * mm, 194 * mm, 13 * mm)
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#65758B"))
        canvas.drawString(16 * mm, 8.5 * mm, "Grupo RS Central • Relatório gerado pelo sistema")
        canvas.drawRightString(194 * mm, 8.5 * mm, f"Página {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(output_path, pagesize=A4, rightMargin=16 * mm, leftMargin=16 * mm, topMargin=13 * mm, bottomMargin=17 * mm, title=payload.get("title", "Relatório de estoque"), author="Grupo RS Central")
    story = []
    logo = Image(logo_path, width=20 * mm, height=20 * mm) if logo_path and os.path.exists(logo_path) else Paragraph("<b>RS</b>", styles["ReportTitle"])
    brand = Paragraph("<font size='15'><b>GRUPO RS CENTRAL</b></font><br/><font color='#F59E0B' size='7'><b>GESTÃO INTELIGENTE DE EQUIPAMENTOS</b></font>", styles["Meta"])
    header = Table([[logo, brand]], colWidths=[24 * mm, 152 * mm], rowHeights=[21 * mm])
    header.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0), ("TOPPADDING", (0, 0), (-1, -1), 0)]))
    story += [header, Spacer(1, 4 * mm), Paragraph(payload.get("title", "Relatório de estoque"), styles["ReportTitle"]), Paragraph(f"Filial: <b>{payload.get('branch', '-')}</b> &nbsp;&nbsp;•&nbsp;&nbsp; Situação: <b>{payload.get('status', '-')}</b> &nbsp;&nbsp;•&nbsp;&nbsp; Gerado em: {payload.get('generated_at', '-')}", styles["Meta"]), Spacer(1, 4 * mm)]

    operators = distributions(rows, "operator")
    missing_chip = sum(1 for item in payload.get("products", []) if any(key in item for key in ("iccid", "chip_number", "numero_chip")) and not str(item.get("iccid") or item.get("chip_number") or item.get("numero_chip") or "").strip())
    cards = [("TOTAL", len(rows), BLUE), ("VIVO", operators.get("Vivo", 0), GREEN), ("CLARO", operators.get("Claro", 0), "E35D6A"), ("SEM CHIP", missing_chip, MUTED)]
    card_cells = []
    for label, value, color in cards:
        drawing = Drawing(12 * mm, 12 * mm)
        drawing.add(Circle(6 * mm, 6 * mm, 5.5 * mm, fillColor=colors.HexColor("#EAF4FA"), strokeColor=None))
        if label == "TOTAL":
            stroke = colors.HexColor("#0B7FC1")
            cube = Path()
            cube.moveTo(3.4 * mm, 7.3 * mm); cube.lineTo(6 * mm, 8.8 * mm); cube.lineTo(8.6 * mm, 7.3 * mm)
            cube.lineTo(8.6 * mm, 4.3 * mm); cube.lineTo(6 * mm, 2.8 * mm); cube.lineTo(3.4 * mm, 4.3 * mm); cube.closePath()
            cube.strokeColor = None
            cube.fillColor = stroke
            drawing.add(cube)
            white = colors.white
            drawing.add(Line(3.8 * mm, 7.1 * mm, 6 * mm, 5.85 * mm, strokeColor=white, strokeWidth=.75))
            drawing.add(Line(8.2 * mm, 7.1 * mm, 6 * mm, 5.85 * mm, strokeColor=white, strokeWidth=.75))
            drawing.add(Line(6 * mm, 5.85 * mm, 6 * mm, 3.25 * mm, strokeColor=white, strokeWidth=.75))
        elif label == "SEM CHIP":
            stroke = colors.HexColor("#7A899D")
            sim = Path()
            sim.moveTo(3.5 * mm, 3 * mm); sim.lineTo(3.5 * mm, 9 * mm); sim.lineTo(7.2 * mm, 9 * mm)
            sim.lineTo(8.5 * mm, 7.7 * mm); sim.lineTo(8.5 * mm, 3 * mm); sim.closePath()
            sim.strokeColor = stroke; sim.strokeWidth = 1.05; sim.fillColor = None
            drawing.add(sim)
            drawing.add(Line(3.1 * mm, 2.7 * mm, 8.9 * mm, 9.3 * mm, strokeColor=stroke, strokeWidth=1.35))
        else:
            stroke = colors.HexColor(f"#{color}")
            bar_width = 1.25 * mm
            for x, height in ((3.0, 2.6), (4.9, 4.2), (6.8, 5.8), (8.7, 7.4)):
                drawing.add(Rect(x * mm, 2.4 * mm, bar_width, height * mm, rx=.45 * mm, ry=.45 * mm, fillColor=stroke, strokeColor=None))
        label_text = Paragraph(f"<font color='#{MUTED}' size='7'>{label.title()}</font>", styles["Meta"])
        value_text = Paragraph(f"<font color='#{color}' size='17'><b>{value}</b></font>", styles["Meta"])
        text = Table([[label_text], [value_text]], colWidths=[25 * mm], rowHeights=[4 * mm, 9 * mm])
        text.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0), ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0)]))
        nested = Table([[drawing, text]], colWidths=[13 * mm, 25 * mm], rowHeights=[14 * mm])
        nested.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0), ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0)]))
        card_cells.append(nested)
    card_table = Table([card_cells], colWidths=[43 * mm] * 4, rowHeights=[18 * mm], hAlign="LEFT")
    card_table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F4F8FC")), ("BOX", (0, 0), (-1, -1), .6, colors.HexColor("#D8E4EF")), ("INNERGRID", (0, 0), (-1, -1), .6, colors.HexColor("#D8E4EF")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("ALIGN", (0, 0), (-1, -1), "CENTER")]))
    story += [card_table]

    models = distributions(rows, "model")
    story.append(Paragraph("Distribuição por modelo", styles["Section"]))
    max_count = max(models.values(), default=1)
    model_data = []
    for name, count in models.most_common(5):
        width = max(2, int(75 * count / max_count))
        bar = Table([[""]], colWidths=[width * mm], rowHeights=[3.2 * mm])
        bar.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#0B7FC1"))]))
        model_data.append([Paragraph(clean(name), styles["Meta"]), bar, Paragraph(f"<b>{count}</b>", styles["Meta"])])
    if not model_data:
        model_data = [[Paragraph("Sem dados", styles["Meta"]), "", "0"]]
    model_table = Table(model_data, colWidths=[38 * mm, 120 * mm, 10 * mm], rowHeights=7 * mm)
    model_table.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 3)]))
    story.append(model_table)
    story.append(Paragraph("Lista de equipamentos", styles["Section"]))

    data = [["NÚMERO DE SÉRIE", "IDENTIFICAÇÃO", "MODELO", "OPERADORA", "SITUAÇÃO", "ATUALIZAÇÃO"]]
    for row in rows:
        data.append([row["serial"], row["plate"], row["model"], row["operator"], row["status"], row["updated"]])
    table = Table(data, repeatRows=1, colWidths=[30 * mm, 25 * mm, 31 * mm, 25 * mm, 27 * mm, 38 * mm], hAlign="LEFT")
    commands = [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0B3154")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, 0), 7), ("FONTNAME", (0, 1), (-1, -1), "Helvetica"), ("FONTSIZE", (0, 1), (-1, -1), 7.2), ("TEXTCOLOR", (0, 1), (-1, -1), colors.HexColor("#17263A")), ("GRID", (0, 0), (-1, -1), .45, colors.HexColor("#D8E4EF")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]
    for index in range(1, len(data)):
        if index % 2 == 0:
            commands.append(("BACKGROUND", (0, index), (-1, index), colors.HexColor("#F4F8FC")))
    table.setStyle(TableStyle(commands))
    story.append(table)
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return {"ok": True, "path": output_path, "format": "pdf", "rows": len(rows)}


def generate_maintenance_pdf(payload, output_path):
    from html import escape
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

    styles = getSampleStyleSheet()
    styles['Heading1'].textColor = colors.HexColor('#173b5d')
    styles['Heading2'].textColor = colors.HexColor('#2879bf')
    rows = payload.get('visits', [])
    states = {'pendente': 'Em análise', 'aguardando_peca': 'Aguardando peça',
              'aguardando_cliente': 'Aguardando cliente', 'concluido': 'Concluída', 'cancelado': 'Cancelada'}
    story = [Paragraph('Grupo RS Central · Manutenções', styles['Heading1']),
             Paragraph(escape(str(payload.get('branch', ''))) + f' · {len(rows)} atendimento(s)', styles['Normal']), Spacer(1, 16)]
    for row in rows:
        story.append(Paragraph(escape(str(row.get('plate', ''))) + ' · ' + escape(str(row.get('client', ''))), styles['Heading2']))
        fields = [('Entrada', 'created_at'), ('Aparelho de chegada', 'serial'),
                  ('Motivo', 'reason'), ('Meio', 'discovery_method'), ('Relato', 'note')]
        if row.get('replacement_serial'):
            fields += [('Aparelho para instalação', 'replacement_serial'), ('Registro da baixa', 'stock_discharge_id')]
        if int(row.get('visit_version', 0)) < 2:
            fields += [('Situação histórica', 'status'), ('Aparelho de saída (histórico)', 'departure_serial'),
                       ('Diagnóstico (histórico)', 'diagnosis'), ('Solução (histórico)', 'solution'), ('Responsável (histórico)', 'technician')]
        for label, key in fields:
            value = str(row.get(key, '') or 'Não informado')
            if key == 'status': value = states.get(value, value)
            story.append(Paragraph('<b>' + label + ':</b> ' + escape(value).replace('\n', '<br/>'), styles['Normal']))
        story.append(Spacer(1, 16))
    SimpleDocTemplate(output_path, pagesize=A4, title='Histórico de manutenções', leftMargin=36, rightMargin=36).build(story)
    return {'ok': True, 'path': output_path, 'format': 'pdf', 'rows': len(rows)}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--format", choices=("pdf", "xlsx"), required=True)
    parser.add_argument("--logo", default="")
    args = parser.parse_args()
    with open(args.input, "r", encoding="utf-8-sig") as handle:
        payload = json.load(handle)
    if payload.get('report_type') == 'maintenance' and args.format == 'pdf':
        result = generate_maintenance_pdf(payload, args.output)
    else:
        result = generate_pdf(payload, args.output, args.logo) if args.format == "pdf" else generate_xlsx(payload, args.output, args.logo)
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False))
        sys.exit(1)
